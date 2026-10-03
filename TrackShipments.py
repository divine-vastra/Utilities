import asyncio
import sys
import openpyxl
from datetime import datetime
from playwright.async_api import async_playwright
import time
import json
import re
import os
from pathlib import Path
from copy import copy
from openpyxl.utils import get_column_letter
import gspread
from gspread.utils import rowcol_to_a1
from whatsapp_CloudAPISender import WhatsAppCloudAPI

def load_env_file(path: Path) -> None:
    """Minimal .env loader (no extra package needed). Lines: KEY=VALUE, '#' comments,
    optional quotes and 'export ' prefix. Real environment variables take priority."""
    if not path.is_file():
        print(f"[INFO] No .env file found at {path} — using system environment variables only")
        return
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        if line.lower().startswith("export "):
            line = line[7:].strip()
        key, value = line.split("=", 1)
        key, value = key.strip(), value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in ("'", '"'):
            value = value[1:-1]
        os.environ.setdefault(key, value)

# .env lives in the same folder as this script (not the current working directory)
load_env_file(Path(__file__).resolve().parent / ".env")


# ─── CONFIG ───────────────────────────────────────────────────────────────────
BACKEND      = "gsheet"          # "gsheet" = Google Sheet, "excel" = local EXCEL_FILE
EXCEL_FILE   = "shipments.xlsx"
POLL_SECONDS = 3600  # 60 minutes between full runs

# Google Sheet backend (same sheet that Updateshipments.py appends to)
SHEET_ID  = "1mFvlvN0wIF9sTIPJrldAoyBfolS1rodzUqG_gDtG1hQ"
SHEET_TAB = "Sheet1"
SERVICE_ACCOUNT_FILE = os.environ.get(
    "GOOGLE_SERVICE_ACCOUNT_FILE", str(Path(__file__).resolve().parent / "service_account.json"))
NCOLS = 19  # columns A..S are loaded into memory

# Excel column positions (1-based)
COL_CUSTOMER_NAME = 1   # A: Customer Name         ← NEW
COL_AWB           = 2   # B: AWB / waybill number  ← shifted
COL_CP_NAME       = 3   # C: Courier Partner       ← shifted
COL_STATUS        = 4   # D: Status description    ← shifted
COL_LOCATION      = 5   # E: Last scan location    ← shifted
COL_REMARK        = 6   # F: Last scan remark      ← shifted
COL_TIMESTAMP     = 7   # G: Time of last scan     ← shifted
COL_LAST_UPDATED  = 8   # H: When script last detected a change ← shifted
COL_EDD           = 9   # I: EDD
COL_TRACK_URL     = 10  # J: Tracking URL
COL_DEST          = 12  # L: Destination
COL_ORDER_TYPE    = 13  # M: Order Type - Prepaid or COD
COL_PHONE         = 15  # O: Customer phone number (10 digits, 91 gets prefixed)
COL_WA_DISPATCHED = 17  # Q: Timestamp when Order-Dispatched WhatsApp was sent
COL_WA_OFD        = 18  # R: Timestamp when Out-for-Delivery WhatsApp was sent
COL_WA_DELIVERED  = 19  # S: Timestamp when Order-Delivered WhatsApp was sent


# ─── WHATSAPP CONFIG ──────────────────────────────────────────────────────────
DRY_RUN           = os.environ.get("DRY_RUN", "false").strip().lower() in ("1", "true", "yes", "on")  # True = only log what would be sent; nothing is sent or marked
WA_TOKEN          = os.environ.get("WA_TOKEN", "")
WA_PHONE_NUMBER_ID = os.environ.get("WA_PHONE_NUMBER_ID", "")
WA_TEMPLATE_NAME  = "out_for_delivery"
WA_TEMPLATE_DISPATCHED = "order_dispatched"
WA_TEMPLATE_DELIVERED  = "order_delivered"
WA_LANG           = "en"
WA_FIXED_CP_ID    = "105"
OFD_STATUS        = "out for delivery"
DELIVERED_STATUS  = "delivered"
ARCHIVE_SHEET     = "Delivered Earlier"  # delivered rows are moved here from Sheet1
# ──────────────────────────────────────────────────────────────────────────────

# def build_url(awb: str, cp_id: str) -> str:
#     return f"https://track.clickpost.in/page?waybill={awb}&cp_id={cp_id}"

async def scrape_via_network_intercept(page, url: str) -> dict | None:
    captured = {}

    async def handle_response(response):
        try:
            resp_url = response.url
            if ("clickpost" in resp_url or "track" in resp_url) and response.status == 200:
                content_type = response.headers.get("content-type", "")
                if "json" in content_type:
                    body = await response.json()
                    if any(k in str(body) for k in ["status", "waybill", "scan", "shipment"]):
                        captured["data"] = body
        except Exception:
            pass

    page.on("response", handle_response)
    await page.goto(url, wait_until="networkidle", timeout=30000)
    await page.wait_for_timeout(3000)

    if not captured.get("data"):
        return None

    return parse_intercepted_json(captured["data"])

def parse_intercepted_json(data: dict) -> dict | None:
    try:
        result = data.get("result", {})
        if result:
            shipment = next(iter(result.values()), {})
            if shipment.get("valid"):
                latest = shipment.get("latest_status", {})
                additional = shipment.get("additional", {})
                return {
                    "status":    latest.get("clickpost_status_description")
                                 or latest.get("status", ""),
                    "location":  latest.get("location", ""),
                    "remark":    latest.get("remark", ""),
                    "timestamp": latest.get("timestamp", ""),
                    "edd":       additional.get("courier_partner_edd", ""),
                }

        status = (data.get("status") or data.get("current_status")
                  or data.get("shipment_status") or "")
        if status:
            return {
                "status":    str(status),
                "location":  str(data.get("location") or data.get("city") or ""),
                "remark":    str(data.get("remark") or data.get("message") or ""),
                "timestamp": str(data.get("timestamp") or data.get("updated_at") or ""),
                "edd":       str(data.get("edd") or data.get("expected_delivery") or ""),
            }
    except Exception as e:
        print(f"    [WARN] JSON parse error: {e}")

    return None

async def scrape_via_html(page) -> dict | None:
    try:
        await page.wait_for_selector(
            "[class*='status'], [class*='clickpost_status_bucket_description'], [class*='track'], [class*='Track']",
            timeout=15000
        )
    except Exception:
        print("    [WARN] Timed out waiting for status element on page")

    html = await page.content()

    selectors_to_try = [
        #("[class*='current-status']",  "status"),
        ("[id*='currentStatus']",   "status"),
        #("[class*='status-Text']",     "status"),
        #("[class*='statusText']",      "status"),
        #("[class*='shipment-status']", "status"),
        #("[class*='tracking-status']", "status"),
        ("[class*='location']",        "location"),
        #("[class*='city']",            "location"),
        ("[class*='remark']",          "remark"),
        #("[class*='message']",         "remark"),
        #("[class*='description']",     "remark"),
        ("[id*='currentTimeStamp']",     "timestamp"),
        ("[class*='expected-date pull-left']",     "edd"),
        
        
    ]

    result = {"status": "", "location": "", "remark": "", "timestamp": "", "edd": ""}

    for selector, field in selectors_to_try:
        if result[field]:
            continue
        try:
            el = await page.query_selector(selector)
            if el:
                text = (await el.inner_text()).strip()
                if text and len(text) < 200:
                    result[field] = text
        except Exception:
            pass

    if not result["status"]:
        patterns = [
            r'"current_status"\s*:\s*"([^"]+)"',
            r'"status"\s*:\s*"([^"]+)"',
            r'"clickpost_status_description"\s*:\s*"([^"]+)"',
            r'status["\s:]+([A-Za-z ]{4,30})',
        ]
        for pattern in patterns:
            match = re.search(pattern, html, re.IGNORECASE)
            if match:
                result["status"] = match.group(1).strip()
                break

    return result if result["status"] else None

async def scrape_shipment(browser, awb: str, track_url: str) -> dict | None:
    #url = build_url(awb, cp_id)
    context = await browser.new_context(
        user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        )
    )
    page = await context.new_page()

    try:
        print(f"    → Fetching: {track_url}")

        data = await scrape_via_network_intercept(page, track_url)

        if data:
            print(f"    ✓ Got data via network intercept")
        else:
            print(f"    → Network intercept empty, falling back to HTML scrape...")
            data = await scrape_via_html(page)
            if data:
                print(f"    ✓ Got data via HTML scrape")
            else:
                print(f"    ✗ Could not extract status — check selectors for this page")

        return data

    except Exception as e:
        print(f"    [ERROR] {e}")
        return None
    finally:
        await context.close()

def normalize_phone(raw) -> str | None:
    """Return 91XXXXXXXXXX or None if the number is not usable."""
    if raw is None:
        return None
    if isinstance(raw, float):
        raw = int(raw)
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 12 and digits.startswith("91"):
        return digits
    digits = digits.lstrip("0")
    if len(digits) == 10:
        return "91" + digits
    return None

def is_ofd(status: str) -> bool:
    return (status or "").strip().lower() == OFD_STATUS

def is_delivered(status: str) -> bool:
    return (status or "").strip().lower() == DELIVERED_STATUS

def send_whatsapp(client, template: str, phone: str, awb: str) -> bool:
    """Send a template (tracking-link button, except order_delivered which has none). Returns True on success."""
    try:
        if template == WA_TEMPLATE_DELIVERED:
            components = []
        else:
            components = [client.build_url_button(index=0, value=f"{awb}&cp_id={WA_FIXED_CP_ID}")]
        client.send_template(phone, template, components, WA_LANG)
        return True
    except Exception as e:
        print(f"    ⚠ WhatsApp failed for {phone} (AWB {awb}): {e}")
        return False

def archive_rows(wb, ws, row_numbers: list[int]) -> None:
    """Cut the given rows (ascending row numbers) from ws and append them, with formatting,
    to the bottom of the ARCHIVE_SHEET tab in the same order. Creates the tab if missing."""
    if not row_numbers:
        return
    max_col = ws.max_column

    if ARCHIVE_SHEET in wb.sheetnames:
        arc = wb[ARCHIVE_SHEET]
    else:
        arc = wb.create_sheet(ARCHIVE_SHEET)
        for c in range(1, max_col + 1):  # new tab: copy header row + column widths
            src, dst = ws.cell(1, c), arc.cell(1, c)
            dst.value = src.value
            if src.has_style:
                dst._style = copy(src._style)
            letter = get_column_letter(c)
            if letter in ws.column_dimensions:
                arc.column_dimensions[letter].width = ws.column_dimensions[letter].width

    for r in row_numbers:  # copy all first, delete afterwards
        target = arc.max_row + 1
        for c in range(1, max_col + 1):
            src, dst = ws.cell(r, c), arc.cell(target, c)
            dst.value = src.value
            if src.has_style:
                dst._style = copy(src._style)
            if src.hyperlink:
                dst.hyperlink = src.hyperlink.target or src.hyperlink.location
        if ws.row_dimensions[r].height:
            arc.row_dimensions[target].height = ws.row_dimensions[r].height

    for r in sorted(row_numbers, reverse=True):  # reverse so earlier row numbers stay valid
        ws.delete_rows(r)

def cell_str(v) -> str:
    """Cell value -> trimmed string ('' for blank; whole floats lose the '.0')."""
    if v is None or v == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()

def _retry(fn, *args, **kwargs):
    """Call a gspread function, waiting and retrying on quota (429) / transient errors."""
    for attempt in range(5):
        try:
            return fn(*args, **kwargs)
        except gspread.exceptions.APIError as e:
            code = e.response.status_code
            if code in (429, 500, 503) and attempt < 4:
                wait = 30 * (attempt + 1)
                print(f"    [WARN] Google Sheets API error {code} — retrying in {wait}s")
                time.sleep(wait)
            else:
                raise

def _pad(rows: list[list]) -> list[list]:
    return [list(r) + [""] * (NCOLS - len(r)) if len(r) < NCOLS else list(r[:NCOLS]) for r in rows]

class ExcelStore:
    """Local Excel backend. Saved after every WhatsApp send (crash-safe)."""
    save_after_send = True

    def __init__(self):
        try:
            self.wb = openpyxl.load_workbook(EXCEL_FILE)
        except Exception as e:
            print(f"ERROR: Could not read '{EXCEL_FILE}' — {e}")
            sys.exit(1)
        self.ws = self.wb["Sheet1"]

    def load(self) -> list[list]:
        rows = [["" if v is None else v for v in r]
                for r in self.ws.iter_rows(min_row=1, max_col=NCOLS, values_only=True)]
        return _pad(rows)

    def flush(self, dirty: dict) -> None:
        for (r, c), v in dirty.items():
            self.ws.cell(row=r, column=c).value = v
        self.wb.save(EXCEL_FILE)

    def archive(self, row_numbers: list[int], awbs: set) -> None:
        archive_rows(self.wb, self.ws, row_numbers)
        self.wb.save(EXCEL_FILE)

class GSheetStore:
    """Google Sheet backend. Changes are written in one batch at the end of the run."""
    save_after_send = False

    def __init__(self):
        # GitHub Actions / CI: JSON key content in env GOOGLE_CREDS_JSON (a secret).
        # Local: key file at SERVICE_ACCOUNT_FILE.
        sa_json = os.environ.get("GOOGLE_CREDS_JSON", "").strip()
        try:
            if sa_json:
                gc = gspread.service_account_from_dict(json.loads(sa_json))
            elif os.path.exists(SERVICE_ACCOUNT_FILE):
                gc = gspread.service_account(filename=SERVICE_ACCOUNT_FILE)
            else:
                print(f"ERROR: No Google credentials — set GOOGLE_CREDS_JSON or provide '{SERVICE_ACCOUNT_FILE}'")
                sys.exit(1)
        except ValueError as e:
            print(f"ERROR: GOOGLE_CREDS_JSON is not valid JSON — {e}")
            sys.exit(1)
        try:
            self.sh = gc.open_by_key(SHEET_ID)
            self.ws = self.sh.worksheet(SHEET_TAB)
        except Exception as e:
            print(f"ERROR: Could not open Google Sheet — {e}")
            sys.exit(1)

    def _values(self) -> list[list]:
        return _retry(self.ws.get_all_values, value_render_option="UNFORMATTED_VALUE")

    def load(self) -> list[list]:
        rows = _pad(self._values())
        return rows or [[""] * NCOLS]

    def flush(self, dirty: dict) -> None:
        by_row: dict[int, list[int]] = {}
        for (r, c) in dirty:
            by_row.setdefault(r, []).append(c)
        data = []
        for r, cols in by_row.items():
            cols.sort()
            start = prev = cols[0]
            run = [start]
            for c in cols[1:] + [None]:
                if c is not None and c == prev + 1:
                    run.append(c)
                    prev = c
                    continue
                data.append({"range": f"{rowcol_to_a1(r, run[0])}:{rowcol_to_a1(r, run[-1])}",
                             "values": [[dirty[(r, k)] for k in run]]})
                if c is not None:
                    run, prev = [c], c
        for i in range(0, len(data), 400):
            _retry(self.ws.batch_update, data[i:i + 400], value_input_option="RAW")

    def archive(self, row_numbers: list[int], awbs: set) -> None:
        """Copy delivered rows (values + formatting) to ARCHIVE_SHEET, then delete them from
        Sheet1, in one atomic request. Rows are re-located by AWB in case the sheet changed
        (e.g. Updateshipments.py appended rows) while the scrape was running."""
        fresh = self._values()
        rows = [i + 1 for i, row in enumerate(fresh)
                if i > 0 and len(row) > 1 and cell_str(row[1]) in awbs]
        if not rows:
            return
        ws = self.ws
        try:
            arc = self.sh.worksheet(ARCHIVE_SHEET)
        except gspread.exceptions.WorksheetNotFound:
            arc = _retry(self.sh.add_worksheet, title=ARCHIVE_SHEET, rows=len(rows) + 100, cols=ws.col_count)
        if arc.col_count < ws.col_count:
            _retry(arc.add_cols, ws.col_count - arc.col_count)

        def copy_req(src_row: int, dst_row: int) -> dict:
            return {"copyPaste": {
                "source": {"sheetId": ws.id, "startRowIndex": src_row - 1, "endRowIndex": src_row,
                           "startColumnIndex": 0, "endColumnIndex": ws.col_count},
                "destination": {"sheetId": arc.id, "startRowIndex": dst_row - 1, "endRowIndex": dst_row,
                                "startColumnIndex": 0, "endColumnIndex": ws.col_count},
                "pasteType": "PASTE_NORMAL"}}

        requests = []
        used = len(_retry(arc.col_values, 1))
        if used == 0:                      # new/empty tab: copy header row first
            requests.append(copy_req(1, 1))
            used = 1
        needed = used + len(rows)
        if arc.row_count < needed:
            _retry(arc.add_rows, needed - arc.row_count)
        for n, r in enumerate(rows, start=1):
            requests.append(copy_req(r, used + n))
        for r in sorted(rows, reverse=True):  # reverse so earlier row numbers stay valid
            requests.append({"deleteDimension": {"range": {
                "sheetId": ws.id, "dimension": "ROWS", "startIndex": r - 1, "endIndex": r}}})
        _retry(self.sh.batch_update, {"requests": requests})

async def run():
    start_time = time.perf_counter()
    store = GSheetStore() if BACKEND == "gsheet" else ExcelStore()
    grid = store.load()
    dirty: dict[tuple[int, int], object] = {}   # (row, col) -> new value, written back via store.flush

    def setc(r: int, c: int, value) -> None:
        grid[r - 1][c - 1] = value
        dirty[(r, c)] = value

    updated_count = 0
    total_rows = 0
    wa_sent_count = 0
    rows_to_archive = []  # delivered rows, moved to ARCHIVE_SHEET after the loop
    archive_awbs = set()

    for col, header in ((COL_WA_OFD, "WA OFD Sent"),
                        (COL_WA_DISPATCHED, "WA Dispatched Sent"),
                        (COL_WA_DELIVERED, "WA Delivered Sent")):
        if not grid[0][col - 1]:
            setc(1, col, header)

    wa_client = None
    if not DRY_RUN:
        if not WA_TOKEN or not WA_PHONE_NUMBER_ID:
            print("ERROR: set WA_TOKEN and WA_PHONE_NUMBER_ID environment variables (or set DRY_RUN = True).")
            sys.exit(1)
        wa_client = WhatsAppCloudAPI(token=WA_TOKEN, phone_number_id=WA_PHONE_NUMBER_ID)
    else:
        print("  *** DRY_RUN is ON — no WhatsApp messages will be sent ***")

    print(f"\n{'='*60}")
    print(f"  ClickPost Page Scraper — {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ({BACKEND})")
    print(f"{'='*60}")

    try:
        async with async_playwright() as pw:
            browser = await pw.chromium.launch(headless=True)

            for r in range(2, len(grid) + 1):  # Row 1 = headers
                row = grid[r - 1]
                customer_name = cell_str(row[COL_CUSTOMER_NAME - 1])
                awb           = cell_str(row[COL_AWB   - 1]) or None
                cp_name       = cell_str(row[COL_CP_NAME - 1]) or None
                dest          = cell_str(row[COL_DEST - 1]) or None
                order_type    = cell_str(row[COL_ORDER_TYPE - 1]) or None
                track_url     = cell_str(row[COL_TRACK_URL - 1]) or None

                if not awb or not cp_name:
                    continue

                total_rows += 1
                print(f"\n{total_rows:02}. → Customer: {customer_name or 'N/A'} | Dest: {dest} | TYPE: {order_type} | CP_NAME: {cp_name}")

                new = await scrape_shipment(browser, awb, track_url)

                if not new:
                    print(f"  Skipping — no data returned")
                    continue

                new["location"] = " ".join(str(new["location"]).split())  # drop newlines / extra whitespace

                old_status = cell_str(row[COL_STATUS - 1])

                if new["status"] and new["status"] != old_status:
                    print(f"    ✅ STATUS CHANGED: [{old_status}] → [{new['status']}], EDD: {new['edd']}, {new['location'][:30]} ")
                    updated_count += 1
                else:
                    print(f"    — No change: {new['status']},  EDD: {new['edd']}, {new['location'][:30]}")

                setc(r, COL_STATUS,       new["status"])
                setc(r, COL_LOCATION,     new["location"])
                setc(r, COL_REMARK,       new["remark"])
                setc(r, COL_TIMESTAMP,    new["timestamp"])
                setc(r, COL_LAST_UPDATED, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                if new["edd"]:
                    setc(r, COL_EDD, new["edd"])

                # ── WhatsApp events (each sent once per shipment, tracked by its own column) ──
                #   blank -> non-blank  : order_dispatched
                #   any   -> Out for Delivery : out_for_delivery
                #   any   -> Delivered  : order_delivered
                events = []
                if not old_status and new["status"].strip():
                    events.append(("Dispatched", WA_TEMPLATE_DISPATCHED, COL_WA_DISPATCHED))
                if is_ofd(new["status"]) and not is_ofd(old_status):
                    events.append(("OFD", WA_TEMPLATE_NAME, COL_WA_OFD))
                if is_delivered(new["status"]) and not is_delivered(old_status):
                    events.append(("Delivered", WA_TEMPLATE_DELIVERED, COL_WA_DELIVERED))

                for label, template, sent_col in events:
                    if row[sent_col - 1]:
                        continue
                    phone = normalize_phone(row[COL_PHONE - 1])
                    if not phone:
                        print(f"    ⚠ {label} but phone invalid/missing in column O — will retry next poll")
                    elif DRY_RUN:
                        button = "no components" if template == WA_TEMPLATE_DELIVERED else f"button value '{awb}&cp_id={WA_FIXED_CP_ID}'"
                        print(f"    [DRY_RUN] Would send '{template}' to {phone} (AWB {awb}, {button})")
                    elif send_whatsapp(wa_client, template, phone, awb):
                        setc(r, sent_col, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                        wa_sent_count += 1
                        print(f"    📲 WhatsApp '{template}' sent to {phone} (AWB {awb})")
                        if store.save_after_send:
                            try:
                                store.flush(dirty)  # persist immediately to avoid duplicates after a crash
                                dirty.clear()
                            except Exception as e:
                                print(f"    [WARN] Could not save after send (is the Excel file open?): {e}")

                # Delivered rows move to the archive tab regardless of WhatsApp outcome (no retry)
                if is_delivered(new["status"]):
                    rows_to_archive.append(r)
                    archive_awbs.add(awb)

                await asyncio.sleep(2)

            await browser.close()
    finally:
        # Always write what we have (also on Ctrl+C / crash) so sent-timestamps aren't lost
        try:
            store.flush(dirty)
            dirty.clear()
        except Exception as e:
            print(f"ERROR: Could not save results — {e}")

    if rows_to_archive:
        if DRY_RUN:
            print(f"  [DRY_RUN] Would move {len(rows_to_archive)} delivered row(s) to '{ARCHIVE_SHEET}'")
        else:
            store.archive(rows_to_archive, archive_awbs)
            print(f"  Moved {len(rows_to_archive)} delivered row(s) to '{ARCHIVE_SHEET}'")

    print(f"\n{'='*60}")
    print(f"  Done. {updated_count}/{total_rows} shipments updated. WhatsApp sent: {wa_sent_count}. Time taken: {(time.perf_counter() - start_time) / 60:.1f} min")
    print(f"  Finished at — {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} ")
        
    print(f"{'='*60}\n")

if __name__ == "__main__":
    while True:
        asyncio.run(run())
        if os.environ.get("GITHUB_ACTIONS") == "true":  # the workflow's schedule is the loop: run once
            break
        print(f"Sleeping {POLL_SECONDS // 60} minutes before next poll...\n")
        time.sleep(POLL_SECONDS)
